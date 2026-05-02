"""Excel 出力。

- シート「公募中一覧」：新規補助金を上に追加。新規行は薄黄色塗り。
  締切が7日以内のセルは薄赤塗り。
- シート「更新履歴」：既存補助金の変更を追記。

新規行の色付けは「今回の実行で追加された行」のみが対象。前回以前の新規分は
通常色に戻す（実行のたびに塗り直す）。
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

from dateutil import parser as date_parser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models import HEADERS, UPDATE_HEADERS, Subsidy

SHEET_MAIN = "公募中一覧"
SHEET_UPDATES = "更新履歴"

FILL_NEW = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
FILL_DEADLINE = PatternFill(start_color="FFF8CECC", end_color="FFF8CECC", fill_type="solid")
FILL_NONE = PatternFill(fill_type=None)
END_DATE_COL = HEADERS.index("締切日") + 1


def _ensure_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    main = wb.active
    main.title = SHEET_MAIN
    main.append(HEADERS)
    _style_header(main)
    updates = wb.create_sheet(SHEET_UPDATES)
    updates.append(UPDATE_HEADERS)
    _style_header(updates)
    return wb


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, _ in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(i)].width = 20


def _parse_date(s: str):
    if not s:
        return None
    try:
        return date_parser.parse(s).date()
    except (ValueError, TypeError, OverflowError):
        return None


def _apply_deadline_color(ws, row_idx: int, end_date_str: str, today: date) -> None:
    end = _parse_date(end_date_str)
    if end is None:
        return
    if today <= end <= today + timedelta(days=7):
        ws.cell(row=row_idx, column=END_DATE_COL).fill = FILL_DEADLINE


def write(path: Path, new_items: Iterable[Subsidy], updates: Iterable[tuple]) -> None:
    """Excel に新規・更新を書き込む。

    new_items: 今回の実行で新規追加された Subsidy のリスト
    updates:   (Subsidy, prev_payload_dict) のタプルリスト
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure_workbook(path)
    main = wb[SHEET_MAIN]
    upd = wb[SHEET_UPDATES]
    today = date.today()

    # 既存「新規」行の塗りをリセット（締切色は最後にまとめて再適用）
    for row in main.iter_rows(min_row=2):
        for cell in row:
            if cell.fill == FILL_NEW:
                cell.fill = FILL_NONE

    # 新規行を 2 行目（ヘッダ直下）に挿入し、上に積む
    new_items = list(new_items)
    if new_items:
        main.insert_rows(2, amount=len(new_items))
        for offset, item in enumerate(new_items):
            row_idx = 2 + offset
            for col_idx, value in enumerate(item.to_row(), start=1):
                cell = main.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = FILL_NEW

    # 全行に対して締切色を再評価
    for row_idx in range(2, main.max_row + 1):
        end_str = main.cell(row=row_idx, column=END_DATE_COL).value
        _apply_deadline_color(main, row_idx, end_str or "", today)

    # 更新履歴
    detected_at = datetime.now().isoformat(timespec="seconds")
    for s, prev in updates:
        upd.append([
            s.name,
            s.detail_url,
            (prev or {}).get("content_hash_pseudo", ""),
            s.content_hash,
            _summarize_payload(prev),
            _summarize_payload(s.to_dict()),
            detected_at,
        ])

    wb.save(path)


def _summarize_payload(payload) -> str:
    if not payload:
        return ""
    keys = ["end_date", "max_amount", "subsidy_rate", "overview"]
    return " / ".join(f"{k}={payload.get(k, '')}" for k in keys)
